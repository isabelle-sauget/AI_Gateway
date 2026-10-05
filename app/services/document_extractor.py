"""Extraction service for converting DOCX and PDF (digital/scanned) into text."""

import io
from pathlib import Path
import docx
import pymupdf
from PIL import Image
import pytesseract

# Optional: If tesseract is not in your Windows PATH, specify it explicitly:
# pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract all paragraph and table text from an in-memory DOCX file."""
    doc = docx.Document(io.BytesIO(file_bytes))
    extracted_chunks: list[str] = []

    # 1. Paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            extracted_chunks.append(text)

    # 2. Tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                extracted_chunks.append(" | ".join(row_text))

    return "\n".join(extracted_chunks)


def extract_text_from_pdf(file_bytes: bytes, ocr_threshold_chars: int = 40) -> str:
    """
    Extract text from PDF using a hybrid pipeline.
    If a page contains fewer characters than `ocr_threshold_chars`,
    fallback to Tesseract OCR for that page.
    """
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    full_text_pages: list[str] = []

    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            # 1. Attempt native digital text extraction
            page_text = page.get_text("text").strip() # type: ignore

            # 2. Fallback to OCR if page has little to no native text (scanned page)
            if len(page_text) < ocr_threshold_chars:
                # Render page to high-res pixmap (300 DPI for reliable OCR accuracy)
                pix = page.get_pixmap(dpi=300)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                
                # Perform OCR with Romanian and English language dictionaries
                ocr_text = pytesseract.image_to_string(img, lang="ron+eng")
                page_text = ocr_text.strip()

            if page_text:
                full_text_pages.append(page_text)

        return "\n\n".join(full_text_pages)
    finally:
        doc.close()


def extract_document_text(filename: str, file_bytes: bytes) -> str:
    """Route document to the appropriate parser based on file extension."""
    ext = Path(filename).suffix.lower()

    if ext == ".docx":
        text = extract_text_from_docx(file_bytes)
    elif ext == ".pdf":
        text = extract_text_from_pdf(file_bytes)
    elif ext == ".txt":
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported file format: {ext}. Only PDF, DOCX, and TXT are supported.")

    if not text.strip():
        raise ValueError("Could not extract any readable text from the uploaded document.")

    return text